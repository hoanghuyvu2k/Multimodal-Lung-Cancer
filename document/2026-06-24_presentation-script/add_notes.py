# -*- coding: utf-8 -*-
"""Thêm speaker notes vào slides.pptx (2026-06-24) từ kich-ban-thuyet-trinh.md."""
import re
from pptx import Presentation

BASE = r"D:\code\master\doan\code\document\2026-06-24_presentation-script"
md = open(BASE + r"\kich-ban-thuyet-trinh.md", encoding="utf-8").read()

# ---- trích lời thoại từng "## Slide N" ----
def clean_bq(block):
    """Lấy các dòng blockquote '> ...' -> text, bỏ dấu ngoặc kép bao ngoài."""
    out = []
    for ln in block.splitlines():
        m = re.match(r"\s*>\s?(.*)$", ln)
        if m is None:
            if out and ln.strip() == "":
                continue
            if out:
                break
            continue
        out.append(m.group(1))
    txt = "\n".join(out).strip()
    txt = re.sub(r"\n{2,}", "\n\n", txt)
    return txt.strip().strip('"').strip()

slide_script = {}
parts = re.split(r"^## Slide (\d+)[^\n]*\n", md, flags=re.M)
# parts = [pre, '1', body1, '2', body2, ...]
for i in range(1, len(parts) - 1, 2):
    n = int(parts[i]); body = parts[i + 1]
    m = re.search(r"🎙️\s*\*\*Lời thuyết trình\*\*[^\n]*\n(.*?)(?:\n---|\Z)", body, flags=re.S)
    if m:
        slide_script[n] = clean_bq(m.group(1))

# ---- trích 🎙️ *Nói:* của 5 chỉ số sống còn (Phụ lục B) ----
noi = {}
for k in range(1, 6):
    m = re.search(rf"###\s*{k}\.\s.*?🎙️\s*\*Nói:\*\s*(.*?)(?:\n###|\n---|\Z)", md, flags=re.S)
    if m:
        noi[k] = m.group(1).strip().strip('"').strip()

# ---- ghép note cho slide chỉ số ----
def join_noi(keys, head):
    return head + "\n\n" + "\n\n".join(f"• {noi[k]}" for k in keys if k in noi)

metric1 = join_noi([1, 2, 5],
    "Ba chỉ số nhóm PHÂN BIỆT / PHÂN TẦNG nguy cơ:")
metric2 = join_noi([3, 4],
    "Hai chỉ số nhóm HỒI QUY NGUY CƠ & HIỆU CHUẨN:")

agenda = ("Đây là bố cục bài trình bày: bắt đầu từ bối cảnh lâm sàng và khoảng trống nghiên cứu; "
    "tiếp đến hai đóng góp mới về cơ chế attention và mã hoá NLP cho biến lâm sàng; "
    "phần kết quả phân loại đáp ứng; rồi phân tích sống còn — phần có tín hiệu mạnh nhất; "
    "và cuối cùng là bàn luận, hạn chế, kết luận.")
thanks = ("Em xin kết thúc phần trình bày ở đây. Em cảm ơn thầy cô và các bạn đã lắng nghe, "
    "và rất mong nhận được câu hỏi cùng góp ý từ hội đồng. "
    "(Xem Phụ lục Q&A trong file kịch bản để chuẩn bị các câu hỏi thường gặp.)")

# ---- map pptx slide (1-based) -> note ----
MAP = {1: slide_script.get(1), 2: agenda}
MAP.update({3+i: slide_script.get(2+i) for i in range(14)})   # pptx 3..16 -> md 2..15
MAP[17] = metric1
MAP[18] = metric2
MAP.update({19: slide_script.get(16), 20: slide_script.get(17), 21: slide_script.get(18)})
MAP[22] = thanks

# ---- note slide 15 (KQ4 sống còn): giải thích rõ 2 biểu đồ + kết luận ----
NOTE15 = (
"Kết quả thứ tư — theo em là quan trọng nhất — đến từ phân tích sống còn. Khác với AUC nhị phân "
"chỉ hỏi 'có đáp ứng hay không', phân tích sống còn hỏi 'bệnh tiến triển KHI NÀO' và tận dụng được "
"cả những bệnh nhân chưa có biến cố (censored). Em đánh giá trên PFS — thời gian sống không tiến triển. "
"Slide có HAI biểu đồ; em xin giải thích cách đọc từng cái và nó cho ra kết luận gì.\n\n"

"■ BIỂU ĐỒ BÊN TRÁI — Forest plot của hồi quy Cox (Hazard Ratio, HR):\n"
"- Mỗi HÀNG là một mô hình; 8 hàng = 4 cách mã hoá lâm sàng (No clinical / +Labs / +NLP raw / +NLP-PCA16) "
"× 2 nhánh bệnh học (IHC-A, IHC-G).\n"
"- Chấm tròn = giá trị HR ước lượng; đường ngang = khoảng tin cậy 95%; trục X là HR theo thang log; "
"đường thẳng đứng tại HR = 1 là mốc 'không ảnh hưởng'. Chấm càng xa về bên phải của 1 → phân tầng nguy cơ "
"càng mạnh; nếu cả khoảng tin cậy nằm bên phải 1 → có ý nghĩa thống kê.\n"
"- KẾT LUẬN từ biểu đồ này: mọi chấm đều nằm ở HR khoảng 4–6 và toàn bộ khoảng tin cậy nằm bên phải số 1 "
"→ risk score của DyAM là yếu tố tiên lượng ĐỘC LẬP, mạnh và có ý nghĩa ở mọi mô hình, ngay cả sau khi "
"đã hiệu chỉnh cho tuổi, ECOG, albumin, dNLR, di căn gan. Và các hàng có NLP có chấm dịch sang phải hơn "
"(HR cao nhất 6,07 ở IHC-A) → NLP NÂNG được HR, còn thêm số thô (+Labs) thì HR gần như không nhúc nhích.\n\n"

"■ BIỂU ĐỒ BÊN PHẢI — Time-dependent AUC (AUC phụ thuộc thời gian):\n"
"- Hai panel IHC-A và IHC-G. Trục X = thời gian (6 / 12 / 18 tháng); trục Y = AUC tại từng mốc "
"(khả năng phân biệt còn sống / đã tiến triển ở thời điểm đó). 4 đường = 4 cách mã hoá lâm sàng.\n"
"- KẾT LUẬN từ biểu đồ này (rõ nhất ở panel IHC-A): tại 6 tháng bốn đường gần như chồng nhau (~0,72) "
"→ ngắn hạn không khác biệt. Nhưng tại 12–18 tháng, hai đường NLP tách VỌT lên trên (~0,74), còn đường "
"+Labs (số thô) tụt xuống THẤP NHẤT (~0,66 ở 18 tháng), thậm chí dưới cả No clinical. "
"→ Số thô THÊM NHIỄU, còn NLP THÊM TÍN HIỆU, và tín hiệu đó là DÀI HẠN.\n\n"

"★ THÔNG ĐIỆP CỐT LÕI của slide: dù AUC nhị phân của NLP có cải thiện nhưng chưa đạt ý nghĩa do cỡ mẫu nhỏ, "
"thì phân tích sống còn lại cho tín hiệu MẠNH và NHẤT QUÁN nhất. Nó gợi ý rằng mã hoá NLP của biến lâm sàng "
"chủ yếu mang thông tin TIÊN LƯỢNG dài hạn — bệnh tiến triển khi nào — hơn là thông tin dự đoán đáp ứng "
"nhị phân ngắn hạn. Đây là một phân biệt quan trọng cho cách đánh giá các phương pháp mã hoá lâm sàng."
)
MAP[15] = NOTE15

# ---- áp vào pptx ----
prs = Presentation(BASE + r"\slides.pptx")
applied = 0
for idx, sld in enumerate(prs.slides, 1):
    note = MAP.get(idx)
    if note:
        sld.notes_slide.notes_text_frame.text = note
        applied += 1
    else:
        print("  ! slide", idx, "chưa có note")

import os
try:
    prs.save(BASE + r"\slides.pptx"); dest = "slides.pptx"
except Exception:
    prs.save(BASE + r"\slides-with-notes.pptx"); dest = "slides-with-notes.pptx (gốc đang mở)"
print(f"Đã ghi note cho {applied}/22 slide -> {dest}")
